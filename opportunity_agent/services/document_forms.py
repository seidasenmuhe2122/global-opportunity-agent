from __future__ import annotations

import io
from pathlib import Path
import requests
from django.core.files.base import ContentFile
from pypdf import PdfReader, PdfWriter
from docx import Document


def token_value(token, profile, application):
    mapping = {
        'full_name': profile.full_name or application.user.get_full_name() or application.user.username,
        'email': application.user.email, 'phone': profile.phone, 'country': profile.current_country,
        'degree': profile.degree, 'education': profile.education, 'experience': profile.work_experience,
        'skills': ', '.join(profile.skills or []), 'languages': ', '.join(profile.languages or []),
        'certifications': ', '.join(profile.certifications or []), 'linkedin_url': profile.linkedin_url,
        'portfolio_url': profile.portfolio_url, 'github_url': profile.github_url,
        'cover_letter': application.cover_letter, 'job_title': application.opportunity.title,
        'organization': application.opportunity.organization,
    }
    return str(mapping.get(token, token) or '')


def download_form(url, timeout=30):
    r=requests.get(url, timeout=timeout, allow_redirects=True, headers={'User-Agent':'OpportunityAgent/1.0'})
    r.raise_for_status()
    return r.content, r.headers.get('content-type','')


def fill_pdf(template_file, field_map, profile, application):
    data = template_file.read()
    reader = PdfReader(io.BytesIO(data))
    writer = PdfWriter()
    for page in reader.pages: writer.add_page(page)
    values = {field: token_value(token, profile, application) for field, token in (field_map or {}).items()}
    if not reader.get_fields():
        raise ValueError('This PDF has no fillable AcroForm fields. It needs manual review or a site-specific document adapter.')
    writer.update_page_form_field_values(writer.pages[0], values, auto_regenerate=True)
    out=io.BytesIO(); writer.write(out); out.seek(0)
    return out.getvalue()


def fill_docx(template_file, profile, application):
    if hasattr(template_file, 'seek'): template_file.seek(0)
    doc = Document(template_file)
    values = {
        '{{full_name}}': token_value('full_name', profile, application), '{{email}}': token_value('email', profile, application),
        '{{phone}}': token_value('phone', profile, application), '{{country}}': token_value('country', profile, application),
        '{{degree}}': token_value('degree', profile, application), '{{education}}': token_value('education', profile, application),
        '{{experience}}': token_value('experience', profile, application), '{{skills}}': token_value('skills', profile, application),
        '{{languages}}': token_value('languages', profile, application), '{{certifications}}': token_value('certifications', profile, application),
        '{{linkedin_url}}': token_value('linkedin_url', profile, application), '{{portfolio_url}}': token_value('portfolio_url', profile, application),
        '{{github_url}}': token_value('github_url', profile, application), '{{cover_letter}}': token_value('cover_letter', profile, application),
        '{{job_title}}': token_value('job_title', profile, application), '{{organization}}': token_value('organization', profile, application),
    }
    def replace_in_paragraph(paragraph):
        for key,val in values.items():
            if key in paragraph.text:
                for run in paragraph.runs:
                    if key in run.text: run.text=run.text.replace(key,val)
                if key in paragraph.text:
                    # fallback for split runs: rebuild plain paragraph text while preserving content
                    text=paragraph.text
                    for k,v in values.items(): text=text.replace(k,v)
                    if paragraph.runs: paragraph.runs[0].text=text
                    for run in paragraph.runs[1:]: run.text=''
    for paragraph in doc.paragraphs: replace_in_paragraph(paragraph)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs: replace_in_paragraph(paragraph)
    out=io.BytesIO(); doc.save(out); out.seek(0); return out.getvalue()
